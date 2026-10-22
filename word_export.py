
import streamlit as st
import tempfile
from pathlib import Path
import re
from docx import Document

TEMPLATE_PATH = Path(__file__).resolve().parent / "angebot.docx"


def export_doc(template_path, output_path, daten):
    """Replace placeholders while retaining paragraph, table and run styling."""
    doc = Document(str(template_path))
    pattern = re.compile(r"\{(" + "|".join(re.escape(k) for k in sorted(daten, key=len, reverse=True)) + r")\}") if daten else None

    def replace_paragraph(paragraph):
        runs = paragraph.runs
        original = [run.text for run in runs]
        full_text = "".join(original)
        positions = [(i, j) for i, value in enumerate(original) for j in range(len(value))]
        for match in reversed(list(pattern.finditer(full_text))):
            start, end = match.span()
            first, left = positions[start]
            last, right = positions[end - 1]
            value = str(daten[match.group(1)])
            if first == last:
                text = runs[first].text
                runs[first].text = text[:left] + value + text[right + 1:]
            else:
                runs[first].text = runs[first].text[:left] + value
                for index in range(first + 1, last):
                    runs[index].text = ""
                runs[last].text = runs[last].text[right + 1:]

    def replace_table(table):
        for row in table.rows:
            for cell in row.cells:
                for paragraph in cell.paragraphs:
                    replace_paragraph(paragraph)
                for nested in cell.tables:
                    replace_table(nested)

    for paragraph in doc.paragraphs:
        replace_paragraph(paragraph)
    for table in doc.tables:
        replace_table(table)
    for section in doc.sections:
        for part in (section.header, section.footer):
            for paragraph in part.paragraphs:
                replace_paragraph(paragraph)
            for table in part.tables:
                replace_table(table)
    doc.save(str(output_path))
    return str(output_path)

def euro_wert(wert):
    return f"{float(wert):,.2f}".replace(".", "X").replace(",", ".").replace("X", ",") + " €"

def export_word_dokument(name, reisepreis, reise_von, reise_bis, zielgebiet, alter_liste, max_alter,
                         rrv_ev_mit, rrv_ev_ohne,
                         rrv_jv_mit, rrv_jv_ohne,
                         spf_mit, spf_ohne,
                         kranken_mit, kranken_ohne,
                         kv_jv_mit, kv_jv_ohne,
                         rus_ev_mit, rus_ev_ohne,
                         rus_jv_mit, rus_jv_ohne):

    daten = {
        "Kundenname": name,
        "Reisedatum": f"{reise_von} – {reise_bis}",
        "Reisepreis": euro_wert(reisepreis),
        "Anzahl": str(len(alter_liste)),
        "Alter": str(max_alter),
        "Reiseziel": zielgebiet,

        "Reiseruecktritt_Einmal_mit_SB": euro_wert(rrv_ev_mit),
        "Reiseruecktritt_Einmal_ohne_SB": euro_wert(rrv_ev_ohne),
        "Reiseruecktritt_Jahres_mit_SB": euro_wert(rrv_jv_mit),
        "Reiseruecktritt_Jahres_ohne_SB": euro_wert(rrv_jv_ohne),
        "Reiseruecktritt_Jahres_Sparfuchs_mit_SB": euro_wert(spf_mit),
        "Reiseruecktritt_Jahres_Sparfuchs_ohne_SB": euro_wert(spf_ohne),

        "Reisekranken_Einmal_mit_SB": euro_wert(kranken_mit),
        "Reisekranken_Einmal_ohne_SB": euro_wert(kranken_ohne),
        "Reisekranken_Jahres_mit_SB": euro_wert(kv_jv_mit),
        "Reisekranken_Jahres_ohne_SB": euro_wert(kv_jv_ohne),

        "RundumSorglos_Einmal_mit_SB": euro_wert(rus_ev_mit),
        "RundumSorglos_Einmal_ohne_SB": euro_wert(rus_ev_ohne),
        "RundumSorglos_Jahres_mit_SB": euro_wert(rus_jv_mit),
        "RundumSorglos_Jahres_ohne_SB": euro_wert(rus_jv_ohne),
    }

    with tempfile.NamedTemporaryFile(delete=False, suffix=".docx") as tmp:
        angebot_path = export_doc(TEMPLATE_PATH, tmp.name, daten)
        with open(angebot_path, "rb") as f:
            st.download_button(
                label="📄 Word-Datei herunterladen",
                data=f,
                file_name=f"angebot_{name}.docx",
                mime="application/vnd.openxmlformats-officedocument.wordprocessingml.document"
            )
